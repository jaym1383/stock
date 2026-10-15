import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import app as stock_app


class SharedStockTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.original_path = stock_app.STOCK_PATH
        stock_app.STOCK_PATH = Path(self.directory.name) / "stock.xlsx"
        self.first_client = stock_app.app.test_client()
        self.second_client = stock_app.app.test_client()

    def tearDown(self):
        stock_app.STOCK_PATH = self.original_path
        self.directory.cleanup()

    @staticmethod
    def workbook_bytes(quantity):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Phones"
        sheet.append(["Model", "Storage", "Color", "Qty"])
        sheet.append(["Test phone", "128GB", "Black", quantity])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def upload(self, content):
        return self.first_client.post(
            "/api/upload-stock",
            data={"file": (BytesIO(content), "stock.xlsx")},
            content_type="multipart/form-data",
            headers={"Origin": "https://jaym1383.github.io"},
        )

    def shared_stock_bytes(self):
        response = self.second_client.get("/api/stock")
        try:
            return response.data
        finally:
            response.close()

    def test_quantity_save_zero_and_stale_version(self):
        uploaded = self.upload(self.workbook_bytes(3))
        metadata = stock_app.stock_metadata()
        payload = {"path": ["Phones", "Test phone", "128GB", "Black"], "quantity": 0, "expected_quantity": 3, "day": stock_app.business_day(), "request_id": "zero-sale-001", "version": uploaded.json["version"]}
        saved = self.first_client.post("/api/stock/quantity", json=payload)
        self.assertEqual(saved.status_code, 200)
        workbook = load_workbook(BytesIO(self.shared_stock_bytes()))
        self.assertEqual(workbook.active["D2"].value, 3)
        self.assertEqual(self.second_client.get("/api/daily-stock").json["products"][0]["quantity"], 0)
        workbook.close()
        self.assertEqual(stock_app.stock_metadata(), metadata)
        payload["request_id"] = "other-client-001"
        self.assertEqual(self.second_client.post("/api/stock/quantity", json=payload).status_code, 409)

    def test_quantity_validation_and_missing_product(self):
        uploaded = self.upload(self.workbook_bytes(3))
        payload = {"path": ["Phones", "Test phone", "128GB", "Black"], "version": uploaded.json["version"], "day": stock_app.business_day(), "expected_quantity": 3, "request_id": "validation-001"}
        for value in [-1, 1.5, True, "2", None]:
            self.assertEqual(self.first_client.post("/api/stock/quantity", json={**payload, "quantity": value}).status_code, 400)
        payload.update(quantity=1, path=["Phones", "Missing", "128GB", "Black"])
        self.assertEqual(self.first_client.post("/api/stock/quantity", json=payload).status_code, 404)
        response = self.second_client.get("/api/stock")
        self.assertEqual(response.headers["X-Stock-Source"], "stock.xlsx")
        self.assertTrue(response.headers["X-Stock-Imported-At"])
        response.close()

    def test_original_source_name_and_duplicate_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Phones"
        sheet.append(["Model", "Storage", "Color", "Qty"])
        sheet.append(["Test phone", "128GB", "Black", 2])
        sheet.append(["Test phone", "128GB", "Black", 3])
        output = BytesIO()
        workbook.save(output)
        response = self.first_client.post("/api/upload-stock", data={
            "file": (BytesIO(output.getvalue()), "stock.xlsx"),
            "source_name": "original-report.csv",
        })
        self.assertEqual(stock_app.stock_metadata()["source"], "original-report.csv")
        saved = self.first_client.post("/api/stock/quantity", json={
            "path": ["Phones", "Test phone", "128GB", "Black"],
            "quantity": 4, "expected_quantity": 5, "day": stock_app.business_day(), "request_id": "duplicate-row-001", "version": response.json["version"],
        })
        self.assertEqual(saved.status_code, 200)
        updated = load_workbook(BytesIO(self.shared_stock_bytes()))
        self.assertEqual(sum(updated.active.cell(row, 4).value for row in [2, 3]), 5)
        self.assertEqual(self.second_client.get("/api/daily-stock").json["events"][0]["delta"], -1)
        updated.close()
        self.assertEqual(self.first_client.post("/api/stock/quantity", json=[1]).status_code, 400)

    def test_daily_idempotency_reset_and_import(self):
        version = self.upload(self.workbook_bytes(1)).json["version"]
        payload = {"path": ["Phones", "Test phone", "128GB", "Black"], "quantity": 0,
                   "expected_quantity": 1, "version": version, "request_id": "daily-test-001", "day": "2026-10-15"}
        original = self.shared_stock_bytes()
        with patch("app.business_day", return_value="2026-10-15"):
            first = self.first_client.post("/api/sell-out", json=payload)
            self.assertEqual(first.status_code, 200)
            self.assertEqual(first.json["products"][0]["quantity"], 0)
            repeated = self.second_client.post("/api/sell-out", json=payload)
            self.assertEqual(repeated.status_code, 200)
            self.assertEqual(len(repeated.json["events"]), 1)
            payload.update(quantity=2, expected_quantity=0, request_id="daily-test-002")
            incoming = self.second_client.post("/api/sell-out", json=payload)
            self.assertEqual(incoming.status_code, 200)
            self.assertEqual(incoming.json["events"][-1]["delta"], 2)
            self.assertEqual(self.shared_stock_bytes(), original)
            new_version = self.upload(self.workbook_bytes(8)).json["version"]
            state = self.second_client.get("/api/daily-stock").json
            self.assertEqual(state["products"], [])
            self.assertEqual(len(state["events"]), 2)
            payload.update(quantity=1, expected_quantity=2, request_id="daily-test-003")
            self.assertEqual(self.first_client.post("/api/sell-out", json=payload).status_code, 409)
        with patch("app.business_day", return_value="2026-10-16"):
            state = self.second_client.get("/api/daily-stock").json
            self.assertEqual(state["events"], [])
            self.assertEqual(state["products"], [])
            payload["version"] = new_version
            self.assertEqual(self.first_client.post("/api/sell-out", json=payload).status_code, 409)

    def test_two_clients_cannot_sell_last_unit_twice(self):
        from concurrent.futures import ThreadPoolExecutor
        version = self.upload(self.workbook_bytes(1)).json["version"]
        def sell(number):
            with stock_app.app.test_client() as client:
                return client.post("/api/sell-out", json={
                    "path": ["Phones", "Test phone", "128GB", "Black"], "quantity": 0,
                    "expected_quantity": 1, "day": stock_app.business_day(), "version": version,
                    "request_id": f"parallel-sale-{number}",
                }).status_code
        with ThreadPoolExecutor(max_workers=2) as workers:
            self.assertEqual(sorted(workers.map(sell, [1, 2])), [200, 409])
        state = self.second_client.get("/api/daily-stock").json
        self.assertEqual(len(state["events"]), 1)
        self.assertEqual(state["products"][0]["quantity"], 0)

    def test_other_client_sees_new_upload_and_version(self):
        self.assertIsNone(self.second_client.get("/api/stock-version").json["version"])
        self.assertEqual(self.second_client.get("/api/stock").status_code, 404)

        first = self.upload(self.workbook_bytes(2))
        self.assertEqual(first.status_code, 200)
        first_version = first.json["version"]
        self.assertEqual(first.headers["Access-Control-Allow-Origin"], "https://jaym1383.github.io")
        self.assertIn("X-Stock-Version", first.headers["Access-Control-Expose-Headers"])
        self.assertEqual(self.second_client.get("/api/stock-version").json["version"], first_version)
        self.assertEqual(self.shared_stock_bytes(), self.workbook_bytes(2))

        second = self.upload(self.workbook_bytes(5))
        self.assertEqual(second.status_code, 200)
        self.assertNotEqual(second.json["version"], first_version)
        self.assertEqual(self.shared_stock_bytes(), self.workbook_bytes(5))

    def test_invalid_upload_keeps_previous_stock(self):
        valid = self.workbook_bytes(3)
        self.assertEqual(self.upload(valid).status_code, 200)
        rejected = self.upload(b"not an xlsx file")
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(self.shared_stock_bytes(), valid)

    def test_pdf_upload_publishes_shared_stock(self):
        workbook = self.workbook_bytes(7)
        rows = [["Phones", "Test phone", "128GB", "Black", 7, "CODE"]]

        def write_mock(_rows, output_path):
            Path(output_path).write_bytes(workbook)

        with patch("pdf_to_stock.extract_records", return_value=[]), \
             patch("pdf_to_stock.build_rows", return_value=rows), \
             patch("pdf_to_stock.write_workbook", side_effect=write_mock):
            response = self.first_client.post(
                "/api/convert",
                data={"file": (BytesIO(b"%PDF-1.4"), "stock.pdf")},
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-Stock-Rows"], "1")
        self.assertEqual(response.headers["X-Stock-Quantity"], "7")
        self.assertEqual(response.headers["X-Stock-Version"], self.second_client.get("/api/stock-version").json["version"])
        self.assertEqual(self.shared_stock_bytes(), workbook)


if __name__ == "__main__":
    unittest.main()
