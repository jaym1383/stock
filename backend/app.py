from pathlib import Path
from io import BytesIO
from tempfile import TemporaryDirectory
import os
import json
from datetime import datetime, timezone
from threading import RLock
from contextlib import contextmanager
from urllib.parse import quote

try:
    import fcntl
except ImportError:
    fcntl = None

STOCK_LOCK = RLock()

from flask import Flask, jsonify, make_response, request, send_file
from openpyxl import load_workbook


app = Flask(__name__)
STOCK_PATH = Path(os.environ.get("STOCK_DATA_DIR", Path(__file__).resolve().parent.parent / "data")) / "stock.xlsx"


@contextmanager
def stock_lock():
    # PythonAnywhere can serve requests in multiple worker processes.
    with STOCK_LOCK:
        STOCK_PATH.parent.mkdir(parents=True, exist_ok=True)
        with STOCK_PATH.with_suffix(".lock").open("a+b") as lock_file:
            if fcntl is not None:
                fcntl.flock(lock_file, fcntl.LOCK_EX)
            try:
                yield
            finally:
                if fcntl is not None:
                    fcntl.flock(lock_file, fcntl.LOCK_UN)

ALLOWED_ORIGINS = {
    "https://jaym1383.github.io",
    "http://localhost:8000",
    "http://127.0.0.1:8000",
}


@app.after_request
def add_cors_headers(response):
    origin = request.headers.get("Origin")
    if origin in ALLOWED_ORIGINS:
        response.headers["Access-Control-Allow-Origin"] = origin
        response.headers["Vary"] = "Origin"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type"
    response.headers["Access-Control-Expose-Headers"] = "X-Stock-Version, X-Stock-Rows, X-Stock-Quantity, X-Stock-Source, X-Stock-Imported-At"
    return response


@app.get("/")
def index():
    return jsonify({"ok": True, "service": "stock pdf converter"})


@app.get("/health")
def health():
    return jsonify({"ok": True})


def stock_version():
    stat = STOCK_PATH.stat()
    return f"{stat.st_mtime_ns}-{stat.st_size}"


def no_store(response):
    response.headers["Cache-Control"] = "no-store"
    return response


def publish_stock(content, source=None):
    STOCK_PATH.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(dir=STOCK_PATH.parent) as tmp:
        staged = Path(tmp) / "stock.xlsx"
        staged.write_bytes(content)
        os.replace(staged, STOCK_PATH)
    if source is not None:
        metadata = {"source": source, "imported_at": datetime.now(timezone.utc).isoformat()}
        STOCK_PATH.with_suffix(".metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
    return stock_version()


def stock_metadata():
    try:
        return json.loads(STOCK_PATH.with_suffix(".metadata.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {"source": STOCK_PATH.name, "imported_at": None}


@app.get("/api/stock-version")
def get_stock_version():
    version = stock_version() if STOCK_PATH.exists() else None
    return no_store(jsonify({"version": version}))


@app.get("/api/stock")
def get_stock():
    with stock_lock():
        if not STOCK_PATH.exists():
            return no_store(jsonify({"detail": "No shared stock has been uploaded yet."})), 404
        content = STOCK_PATH.read_bytes()
        version = stock_version()
        metadata = stock_metadata()
    response = send_file(BytesIO(content), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response.headers["X-Stock-Version"] = version
    response.headers["X-Stock-Source"] = quote(metadata["source"])
    response.headers["X-Stock-Imported-At"] = metadata.get("imported_at") or ""
    return no_store(response)


@app.route("/api/stock/quantity", methods=["POST", "OPTIONS"])
def update_quantity():
    if request.method == "OPTIONS":
        return make_response("", 204)
    data = request.get_json(silent=True) or {}
    if not isinstance(data, dict):
        return jsonify({"detail": "Invalid request."}), 400
    keys = data.get("path")
    qty = data.get("quantity")
    if not isinstance(keys, list) or len(keys) != 4 or not all(isinstance(k, str) and k for k in keys) or type(qty) is not int or qty < 0:
        return jsonify({"detail": "Invalid product or quantity."}), 400
    with stock_lock():
        if not STOCK_PATH.exists():
            return jsonify({"detail": "Stock not found."}), 404
        if data.get("version") != stock_version():
            return jsonify({"detail": "ข้อมูลถูกเปลี่ยนแล้ว กรุณาโหลดสต็อกล่าสุดแล้วลองอีกครั้ง"}), 409
        workbook = load_workbook(STOCK_PATH)
        matches = []
        for sheet in workbook.worksheets:
            if sheet.title.strip().lower() in {"คำแนะนำ", "instructions", "instruction", "readme", "info"}:
                continue
            headers = {str(c.value or "").strip().lower(): c.column for c in sheet[1]}
            if not {"model", "storage", "color", "qty"}.issubset(headers):
                continue
            for row in sheet.iter_rows(min_row=2):
                def value(name):
                    return str(row[headers[name]-1].value or "").strip()
                key = [value("category") if "category" in headers else sheet.title.strip(), value("model"), value("storage"), value("color")]
                if key == keys:
                    matches.append(row[headers["qty"]-1])
        if not matches:
            workbook.close()
            return jsonify({"detail": "Product not found."}), 404
        matches[0].value = qty
        for cell in matches[1:]:
            cell.value = 0
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        version = publish_stock(output.getvalue())
    return no_store(jsonify({"success": True, "quantity": qty, "version": version}))


@app.route("/api/upload-stock", methods=["POST", "OPTIONS"])
def upload_stock():
    if request.method == "OPTIONS":
        return make_response("", 204)

    uploaded = request.files.get("file")
    if uploaded is None or not (uploaded.filename or "").lower().endswith(".xlsx"):
        return jsonify({"detail": "Upload an XLSX workbook."}), 400

    content = uploaded.read()
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        usable = 0
        for sheet in workbook.worksheets:
            if sheet.title.strip().lower() in {"คำแนะนำ", "instructions", "instruction", "readme", "info"}:
                continue
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                continue
            columns = {str(value or "").strip().lower(): i for i, value in enumerate(headers)}
            if not {"model", "storage", "color", "qty"}.issubset(columns):
                continue
            for row in values:
                try:
                    category = row[columns["category"]] if "category" in columns else sheet.title
                    required = [category, row[columns["model"]], row[columns["storage"]], row[columns["color"]]]
                    quantity = float(row[columns["qty"]])
                    if all(str(value or "").strip() for value in required) and quantity > 0:
                        usable += 1
                except (IndexError, TypeError, ValueError):
                    continue
        workbook.close()
        if not usable:
            return jsonify({"detail": "No usable stock rows found."}), 400
    except Exception:
        return jsonify({"detail": "Invalid XLSX workbook."}), 400

    with stock_lock():
        version = publish_stock(content, request.form.get("source_name") or uploaded.filename)
    return no_store(jsonify({"success": True, "version": version, "rows": usable}))


@app.route("/api/convert", methods=["POST", "OPTIONS"])
def convert_pdf():
    if request.method == "OPTIONS":
        return make_response("", 204)

    uploaded = request.files.get("file")
    if uploaded is None or uploaded.filename == "":
        return jsonify({"detail": "No PDF file was uploaded."}), 400

    filename = uploaded.filename or ""
    content_type = uploaded.content_type or ""
    if content_type != "application/pdf" and not filename.lower().endswith(".pdf"):
        return jsonify({"detail": "Only PDF files are supported."}), 400

    try:
        from pdf_to_stock import build_rows, extract_records, write_workbook
    except SystemExit as exc:
        return jsonify({"detail": str(exc) or "Converter dependency is missing."}), 500
    except ImportError as exc:
        return jsonify({"detail": f"Converter dependency is missing: {exc}"}), 500

    try:
        with TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            pdf_path = tmp_dir / "input.pdf"
            xlsx_path = tmp_dir / "stock.xlsx"

            uploaded.save(pdf_path)
            records = extract_records(str(pdf_path))
            rows = build_rows(records)
            if not rows:
                return jsonify({"detail": "No usable stock rows found in the PDF."}), 400
            write_workbook(rows, str(xlsx_path))
            output = BytesIO(xlsx_path.read_bytes())
            output.seek(0)
            with stock_lock():
                version = publish_stock(output.getvalue(), filename)

            response = send_file(
                output,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name="stock.xlsx",
            )
            response.headers["X-Stock-Rows"] = str(len(rows))
            response.headers["X-Stock-Quantity"] = str(sum(row[4] for row in rows))
            response.headers["X-Stock-Version"] = version
            return no_store(response)
    except Exception as exc:
        return jsonify({"detail": f"Failed to process PDF: {exc}"}), 500


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
